from io import BytesIO

from openpyxl import Workbook
import pytest

from app.templates.excel_import import import_excel_template
from app.templates.service import TemplateService


def workbook_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "位移 Δx"
    sheet["A1"] = "时间 t (s)"
    sheet["B1"] = "速度 v₀\n(m/s)"
    sheet["C1"] = "结果"
    sheet["A2"] = 0
    sheet["B2"] = "α β Ω <保留>"
    sheet["C2"] = "=A2+1"
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "合并说明"
    sheet["C3"] = None
    sheet.freeze_panes = "A2"
    other = book.create_sheet("第二张表")
    other["A1"] = "测量值"
    other["A2"] = "预留"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_excel_import_keeps_visible_sheet_dimensions_text_and_merges() -> None:
    schema, warnings = import_excel_template(workbook_bytes(), "实验模板.xlsx")
    assert len(schema["tables"]) == 2
    first = schema["tables"][0]
    assert (first["grid"]["rows"], first["grid"]["columns"]) == (3, 3)
    cells = {(cell["row"], cell["column"]): cell for cell in first["grid"]["cells"]}
    assert cells[(0, 1)]["text"] == "速度 v₀\n(m/s)"
    assert cells[(1, 0)]["text"] == "0"
    assert cells[(1, 1)]["text"] == "α β Ω <保留>"
    assert cells[(1, 2)]["text"] == "=A2+1"
    assert (cells[(2, 0)]["rowSpan"], cells[(2, 0)]["colSpan"]) == (1, 2)
    assert cells[(2, 2)]["kind"] == "field"
    assert first["grid"]["frozenRows"] == 1
    assert TemplateService().validate(schema) == []
    assert any("公式" in warning for warning in warnings)


def test_excel_import_rejects_oversize_without_truncating() -> None:
    book = Workbook()
    sheet = book.active
    sheet["AE2"] = "超出 30 列"
    output = BytesIO()
    book.save(output)
    with pytest.raises(ValueError, match="未作裁剪"):
        import_excel_template(output.getvalue(), "过宽.xlsx")
