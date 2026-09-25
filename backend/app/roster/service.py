from csv import DictReader
from datetime import UTC, datetime
from io import BytesIO, StringIO
import secrets
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth.errors import AuthError
from app.auth.service import AuthService
from app.db.models import ClassMember, LabClass, RosterImport, User


HEADER_ALIASES = {
    "学号": "student_number",
    "姓名": "name",
    "student_number": "student_number",
    "name": "name",
}
REQUIRED_FIELDS = ("student_number", "name")


class RosterService:
    """Owns parsing, validation and transactional roster commits behind two operations."""

    def __init__(self, auth: AuthService) -> None:
        self.auth = auth

    @staticmethod
    def _now() -> datetime:
        return datetime.now(UTC)

    @staticmethod
    def _read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
        suffix = filename.lower().rsplit(".", 1)[-1]
        if suffix == "csv":
            text = content.decode("utf-8-sig")
            return list(DictReader(StringIO(text)))
        if suffix == "xlsx":
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                return []
            headers = [str(value).strip() if value is not None else "" for value in values[0]]
            return [dict(zip(headers, row, strict=False)) for row in values[1:] if any(value not in (None, "") for value in row)]
        raise AuthError("unsupported_file", "仅支持 .xlsx 或 .csv 名单文件。", 422)

    @staticmethod
    def _normalize(raw: dict[str, Any]) -> dict[str, str]:
        normalized: dict[str, str] = {}
        for header, value in raw.items():
            key = HEADER_ALIASES.get(str(header).strip())
            if key:
                normalized[key] = "" if value is None else str(value).strip()
        return normalized

    def preview(self, db: Session, lab_class: LabClass, actor: User, filename: str, content: bytes) -> RosterImport:
        if len(content) > 5 * 1024 * 1024:
            raise AuthError("file_too_large", "名单文件不能超过 5 MB。", 413)
        rows = [self._normalize(row) for row in self._read_rows(filename, content)]
        if len(rows) > 2000:
            raise AuthError("too_many_rows", "单次名单不能超过 2000 行。", 422)
        errors: list[dict[str, Any]] = []
        seen: set[str] = set()
        valid_rows: list[dict[str, str]] = []
        for index, row in enumerate(rows, start=2):
            for field in REQUIRED_FIELDS:
                if not row.get(field):
                    errors.append({"row": index, "field": field, "message": "该项不能为空。"})
            student_number = row.get("student_number", "")
            if student_number and (len(student_number) > 64 or not student_number.replace("-", "").isalnum()):
                errors.append({"row": index, "field": "student_number", "message": "学号格式不合法。"})
            if student_number in seen:
                errors.append({"row": index, "field": "student_number", "message": "文件内学号重复。"})
            seen.add(student_number)
            existing = db.scalar(select(User).where(User.student_number == student_number)) if student_number else None
            if existing and existing.name != row.get("name"):
                errors.append({"row": index, "field": "name", "message": "数据库已有相同学号但姓名不同，需人工确认。"})
            valid_rows.append(row)
        if not rows:
            errors.append({"row": 1, "field": "file", "message": "文件没有可导入的数据行。"})
        preview = RosterImport(
            class_id=lab_class.id,
            created_by=actor.id,
            source_name=filename[:255],
            status="preview",
            rows_json=valid_rows,
            errors_json=errors,
            created_at=self._now(),
        )
        db.add(preview)
        db.commit()
        return preview

    def commit(self, db: Session, preview: RosterImport) -> dict[str, Any]:
        if preview.status == "committed":
            raise AuthError("already_committed", "该名单已确认导入，不能重复提交。", 409)
        if preview.errors_json:
            raise AuthError("preview_has_errors", "请修正名单错误后重新预览。", 409)
        created = existing_count = memberships_added = 0
        now = self._now()
        try:
            for row in preview.rows_json:
                student_number = row["student_number"]
                user = db.scalar(select(User).where(User.student_number == student_number))
                if user:
                    existing_count += 1
                else:
                    user = User(
                        account=student_number.casefold(),
                        password_hash=self.auth.hash_password(secrets.token_urlsafe(48)),
                        name=row["name"],
                        student_number=student_number,
                        role="student",
                        status="active",
                        must_change_password=False,
                        enrollment_completed=False,
                        created_at=now,
                    )
                    db.add(user)
                    db.flush()
                    created += 1
                membership = db.scalar(
                    select(ClassMember).where(ClassMember.class_id == preview.class_id, ClassMember.user_id == user.id)
                )
                if not membership:
                    db.add(ClassMember(class_id=preview.class_id, user_id=user.id, joined_at=now))
                    memberships_added += 1
            preview.status = "committed"
            preview.committed_at = now
            db.commit()
        except Exception:
            db.rollback()
            raise
        return {
            "id": preview.id,
            "created": created,
            "existing": existing_count,
            "membershipsAdded": memberships_added,
            "requiresEnrollment": created,
        }
