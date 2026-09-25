"""Import an Excel worksheet layout as a reviewable template draft."""

from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


MAX_XLSX_BYTES = 8 * 1024 * 1024


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def import_excel_template(content: bytes, filename: str) -> tuple[dict, list[str]]:
    if not filename.lower().endswith(".xlsx") or not content.startswith(b"PK"):
        raise ValueError("请选择有效的 .xlsx 文件；旧版 .xls 请先另存为 .xlsx。")
    if len(content) > MAX_XLSX_BYTES:
        raise ValueError("Excel 文件不能超过 8 MB。")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(entry.file_size for entry in entries) > 40 * 1024 * 1024:
                raise ValueError("Excel 文件解压后过大，无法安全导入。")
        workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Excel 文件无法读取，请检查文件是否损坏或加密。") from exc

    sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    if not sheets:
        raise ValueError("Excel 文件没有可见工作表。")
    tables = []
    warnings = []
    if len(sheets) != len(workbook.worksheets):
        warnings.append("隐藏工作表未导入；请先在 Excel 中取消隐藏需要的表格。")
    for index, sheet in enumerate(sheets, start=1):
        rows, columns = sheet.max_row, sheet.max_column
        if rows < 2 or rows > 201 or columns < 1 or columns > 30:
            raise ValueError(f"工作表“{sheet.title}”有 {rows} 行、{columns} 列；当前支持 2–201 行、1–30 列，未作裁剪。")
        fields = []
        for column in range(1, columns + 1):
            heading = _cell_text(sheet.cell(1, column).value)
            fields.append({
                "key": f"excel_table_{index}_field_{column}", "label": heading or f"第 {column} 列",
                "symbol": "", "type": "number", "unit": "1", "precision": 2,
                "required": False, "scope": "individual",
            })
        merged = {}
        covered = set()
        for region in sheet.merged_cells.ranges:
            anchor = (region.min_row, region.min_col)
            merged[anchor] = (region.max_row - region.min_row + 1, region.max_col - region.min_col + 1)
            for row in range(region.min_row, region.max_row + 1):
                for column in range(region.min_col, region.max_col + 1):
                    if (row, column) != anchor:
                        covered.add((row, column))
        cells = []
        formulas = False
        for row in range(1, rows + 1):
            for column in range(1, columns + 1):
                if (row, column) in covered:
                    continue
                source = sheet.cell(row, column)
                text = _cell_text(source.value)
                row_span, col_span = merged.get((row, column), (1, 1))
                is_input = row > 1 and not text and row_span == 1 and col_span == 1
                cell = {
                    "id": f"r{row}c{column}", "row": row - 1, "column": column - 1,
                    "rowSpan": row_span, "colSpan": col_span,
                    "kind": "field" if is_input else "label", "text": text,
                }
                if is_input:
                    cell["fieldKey"] = fields[column - 1]["key"]
                cells.append(cell)
                formulas |= source.data_type == "f"
        freeze = sheet.freeze_panes
        frozen_rows = 0
        if isinstance(freeze, str):
            digits = "".join(character for character in freeze if character.isdigit())
            frozen_rows = max(0, min(rows, int(digits) - 1)) if digits else 0
        elif freeze is not None:
            frozen_rows = max(0, min(rows, freeze.row - 1))
        widths = []
        for column in range(1, columns + 1):
            letter = get_column_letter(column)
            excel_width = sheet.column_dimensions[letter].width or 13
            widths.append(max(80, min(320, round(excel_width * 7 + 10))))
        table_id = f"excel_table_{index}"
        tables.append({
            "id": table_id, "title": sheet.title, "description": "",
            "minimumRows": rows - 1, "maximumRows": rows - 1, "fields": fields,
            "grid": {"rows": rows, "columns": columns, "columnWidths": widths,
                     "frozenRows": frozen_rows, "cells": cells},
        })
        if formulas:
            warnings.append(f"“{sheet.title}”中的公式已保留为文字，不会在模板中自动计算。")
    warnings.append("已保留可见工作表的行列、合并区域和单元格内容；数字显示格式、字体、颜色、图片与图表未导入。请核对数据格用途、测量符号、单位及录入类型。")
    return {
        "title": filename[:-5][:160], "description": "", "objectives": [],
        "preparationQuestions": [], "discussionQuestions": [], "steps": [],
        "tables": tables, "calculations": [], "charts": [],
        "checks": [{"type": "minimum_rows", "tableId": table["id"], "value": table["minimumRows"]} for table in tables],
        "requiredArtifacts": [],
    }, warnings
