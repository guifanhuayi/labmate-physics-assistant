"""Convert a DOCX report into an editable, deliberately unapproved template draft."""

from io import BytesIO
import re
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document


MAX_DOCX_BYTES = 8 * 1024 * 1024
_TABLE_TITLE = re.compile(r"^表\s*\d+[.．、:]?\s*", re.I)
_METADATA = {"姓名", "学号", "班级", "教师签字", "分组号"}


def _text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _experiment_title(document: Any, fallback: str) -> str:
    for table in document.tables:
        for row in table.rows:
            values = [_text(cell.text) for cell in row.cells]
            for index, value in enumerate(values[:-1]):
                if value == "实验名称":
                    candidate = next((next_value for next_value in values[index + 1:] if next_value and next_value != value), "")
                    if candidate:
                        return candidate[:160]
    return fallback[:160]


def _grid_and_fields(table: Any, table_index: int) -> tuple[dict[str, Any], list[dict[str, Any]]] | None:
    row_count, column_count = len(table.rows), len(table.columns)
    if row_count < 2 or row_count > 201 or column_count < 2 or column_count > 30:
        return None
    rows = [[cell for cell in row.cells] for row in table.rows]
    first_text = _text(rows[0][0].text)
    if "实验报告附件" in first_text or any(_text(cell.text) in _METADATA for row in rows[:3] for cell in row):
        return None
    if not any(not _text(cell.text) for row in rows[1:] for cell in row):
        return None

    positions: dict[int, list[tuple[int, int]]] = {}
    originals: dict[int, Any] = {}
    for row_index, row in enumerate(rows):
        for column_index, cell in enumerate(row):
            marker = id(cell._tc)
            positions.setdefault(marker, []).append((row_index, column_index))
            originals[marker] = cell

    columns_with_input = {column for row in rows[1:] for column, cell in enumerate(row) if not _text(cell.text)}
    if not columns_with_input:
        return None
    fields = []
    field_keys = {}
    for column in sorted(columns_with_input):
        label = next((
            _text(rows[row][column].text)
            for row in range(1, row_count)
            if _text(rows[row][column].text) and not _text(rows[row][column].text).isdigit()
        ), f"第 {column + 1} 列")
        key = f"word_table_{table_index}_field_{column + 1}"
        field_keys[column] = key
        fields.append({
            "key": key, "label": label[:120], "symbol": "", "type": "number", "unit": "1",
            "precision": 2, "required": False, "scope": "individual",
        })

    cells = []
    for marker, coordinates in sorted(positions.items(), key=lambda entry: min(entry[1])):
        top = min(row for row, _ in coordinates)
        left = min(column for _, column in coordinates)
        bottom = max(row for row, _ in coordinates)
        right = max(column for _, column in coordinates)
        if len(coordinates) != (bottom - top + 1) * (right - left + 1):
            return None
        content = _text(originals[marker].text)
        is_field = not content and top > 0 and left in field_keys and len(coordinates) == 1
        cell = {
            "id": f"r{top + 1}c{left + 1}", "row": top, "column": left,
            "rowSpan": bottom - top + 1, "colSpan": right - left + 1,
            "kind": "field" if is_field else "label", "text": content,
        }
        if is_field:
            cell["fieldKey"] = field_keys[left]
        cells.append(cell)

    title = _TABLE_TITLE.sub("", _text(rows[0][0].text.splitlines()[0]))
    title = title or f"数据记录表 {table_index}"
    data_rows = row_count - 1
    return ({
        "id": f"word_table_{table_index}", "title": title[:160], "description": "",
        "minimumRows": data_rows, "maximumRows": data_rows, "fields": fields,
        "grid": {"rows": row_count, "columns": column_count, "columnWidths": [140] * column_count,
                 "frozenRows": 1, "cells": cells},
    }, fields)


def import_word_template(content: bytes, filename: str) -> tuple[dict[str, Any], list[str]]:
    if not filename.lower().endswith(".docx") or not content.startswith(b"PK"):
        raise ValueError("请选择有效的 .docx Word 文档；旧版 .doc 请先另存为 .docx。")
    if len(content) > MAX_DOCX_BYTES:
        raise ValueError("Word 文档不能超过 8 MB。")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(entry.file_size for entry in entries) > 40 * 1024 * 1024:
                raise ValueError("Word 文档解压后过大，无法安全导入。")
        document = Document(BytesIO(content))
    except ValueError:
        raise
    except BadZipFile as exc:
        raise ValueError("Word 文档无法读取，请检查文件是否损坏或加密。") from exc
    except Exception as exc:
        raise ValueError("Word 文档无法读取，请检查文件是否损坏或加密。") from exc

    fallback = filename.rsplit(".", 1)[0]
    title = _experiment_title(document, fallback)
    warnings = []
    tables = []
    for index, source in enumerate(document.tables, start=1):
        converted = _grid_and_fields(source, index)
        if converted:
            tables.append(converted[0])
    if not tables:
        raise ValueError("未找到可转换的数据记录表；请使用包含空白记录单元格的 Word 表格。")
    if len(tables) < len(document.tables):
        warnings.append(f"已转换 {len(tables)} 张数据表；封面信息、签字栏、图表占位或不规则表格未导入。")

    preparation, discussion, steps = [], [], []
    discussion_section = False
    for paragraph in document.paragraphs:
        line = _text(paragraph.text)
        if not line:
            continue
        if "讨论问题" in line and len(line) < 30:
            discussion_section = True
            continue
        if paragraph.style.name.startswith("List") and ("？" in line or "?" in line):
            (discussion if discussion_section else preparation).append(line[:1000])
        elif not discussion_section and 8 <= len(line) <= 300 and len(steps) < 50:
            steps.append({"id": f"word_step_{len(steps) + 1}", "title": line[:60], "instruction": line})

    if document.inline_shapes:
        warnings.append("文档中的图片未转成操作步骤或测量数据，请人工核对。")
    warnings.append("测量数据的类型、单位、符号、必填要求及公式/图表关系需要教师核对后再保存发布。")
    schema = {
        "title": title, "description": "", "objectives": [],
        "preparationQuestions": preparation[:100], "discussionQuestions": discussion[:100],
        "steps": steps, "tables": tables, "calculations": [], "charts": [],
        "checks": [{"type": "minimum_rows", "tableId": table["id"], "value": table["minimumRows"]} for table in tables],
        "requiredArtifacts": [],
    }
    return schema, warnings
